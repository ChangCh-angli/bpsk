#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 Word 版《BPSK 信号源波形质量报告》(docs/ 下)。

复用 gen_design_doc.py 的样式助手;图片取自 docs/figs/(analyze_wave.py 产出)。
数据更新后先跑 analyze_wave.py 再跑本脚本即可再生成。
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_design_doc import ROOT, set_fonts, add_table  # noqa: E402

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(ROOT, "docs", "波形质量报告_V1.0.docx")
FIGDIR = os.path.join(ROOT, "docs", "figs")


def figure(doc, name, width_cm=16.0):
    path = os.path.join(FIGDIR, name)
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for a in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, a, Cm(2.2))

    set_fonts(doc.styles["Normal"], "Calibri", "宋体", size=11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        set_fonts(doc.styles[name], "Arial", "黑体", size=size, bold=True,
                  color=RGBColor(0, 0, 0))

    # ---- 标题与报告信息 ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BPSK 信号源波形质量报告")
    r.font.size = Pt(20)
    r.bold = True
    r.font.name = "Arial"
    r.element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:eastAsia"), "黑体")
    r.element.rPr.append(rf)

    add_table(doc,
              ["项目", "内容"],
              [["被测对象", "rtl/bpsk_src_top.v(PN23 + RRC α=0.35,13-bit 输出,幅度 = 2797)"],
               ["数据来源", "iverilog RTL 仿真实测导出(tb_bpsk_src.v +dump=),2 000 000 采样 / 163 199 码元,导出同时通过 golden 逐位比对(PASS)"],
               ["分析工具", "scripts/analyze_wave.py(numpy / scipy.signal.welch / matplotlib)"],
               ["日期 / 版本", "2026-09-12 / V1.0"]])

    # ---- 1 时域 ----
    doc.add_heading("1 时域波形", level=1)
    doc.add_paragraph(
        "下图给出 500 个采样(≈41 个码元)的 RTL 输出。波形在 ±A(±2797,绿色点线)附近"
        "游走、码元过渡处过冲,最大值 4094 距满量程 ±4095(红色虚线)仅 1 LSB——这是幅度"
        "定标刻意锁定的零削顶工作点。灰色竖线为码元边界,边界处波形连续、无跳变毛刺。"
    )
    figure(doc, "fig1_time.png")

    # ---- 2 频谱 ----
    doc.add_heading("2 频谱(FFT / Welch 功率谱)", level=1)
    figure(doc, "fig2_spectrum.png")
    add_table(doc,
              ["指标", "实测", "理论", "结论"],
              [["占用带宽 (1+α)Rs", "5.508 MHz", "5.508 MHz", "精确一致"],
               ["主瓣 PSD 高度", "≈ −72 dBFS/Hz", "−71.5 dBFS/Hz", "一致"],
               ["噪声底(12~25 MHz 中位)", "−148.2 dBFS/Hz", "—", "远低于信号,无杂散岛"],
               ["带外/带内功率比", "−23.6 dB", "—", "由 ±4 码元截断主导,见第 3 节"]],
              widths=[5.0, 4.0, 3.6, 3.4])
    doc.add_paragraph("谱形为标准 RRC 主瓣 + 截断形成的 sinc 型旁瓣,平滑滚降、无谱线、无杂散。")

    # ---- 3 EVM ----
    doc.add_heading("3 解调与 EVM", level=1)
    doc.add_paragraph(
        "方法:接收端用理想浮点 RRC 匹配滤波,对采样相位做整拍扫描对齐(最优相位即码元边界,"
        "匹配滤波群延迟 98 = 49+49),符号判决后与 PN23 参照序列求误差。"
    )
    figure(doc, "fig3_evm.png")
    add_table(doc,
              ["指标", "数值"],
              [["EVM_RMS", "4.9236%(−26.2 dB)"],
               ["EVM_peak", "13.308%"],
               ["定时敏感度", "偏移 ±1 采样 → EVM 恶化到约 10.5%(判决时刻必须压在码元边界上)"]],
              widths=[5.0, 11.0])
    doc.add_paragraph("误差分解(符号级直接作差;合成校验 4.9237% vs 实测 4.9236%,闭合):")
    add_table(doc,
              ["误差源", "EVM", "折合 dB"],
              [["RRC ±4 码元截断(ISI)", "4.92e-2", "−26.2 dB(主导)"],
               ["抽头 Q1.14 量化", "7.88e-6", "−102.1 dB"],
               ["13-bit 输出量化 + 取整", "3.03e-5", "−90.4 dB"]],
              widths=[7.0, 4.0, 5.0])
    doc.add_paragraph(
        "结论:全部误差几乎都来自 RRC 成形器的 ±4 码元截断,两处定点量化(−102 / −90 dB)"
        "完全可以忽略。若应用要求更高调制质量,把成形器跨度扩到 ±5 码元(123 阶,结构不变)"
        "即可显著改善。"
    )

    # ---- 4 CCDF ----
    doc.add_heading("4 CCDF 峰均比统计", level=1)
    figure(doc, "fig4_ccdf.png", width_cm=13.5)
    bullet(doc, "RMS = 2552.39 LSB;实测峰值 4094 → 峰均比 4.10 dB(2M 采样内最大值,与全周期 golden 实测一致)")
    bullet(doc, "满量程线 4.11 dB:峰值距削顶仅 1 LSB,零削顶——CCDF 曲线恰好终止于满量程线之前,说明幅度定标既无浪费也无削顶")

    # ---- 5 眼图 ----
    doc.add_heading("5 眼图", level=1)
    figure(doc, "fig5_eye.png", width_cm=14.5)
    doc.add_paragraph(
        "匹配滤波后 1200 码元迹线叠加:交叉点汇聚干净,判决时刻(红色竖线)眼图完全张开,"
        "无可见抖动/噪声模糊,与 EVM 分解结论(量化贡献可忽略)互相印证。"
    )

    # ---- 6 结论 ----
    doc.add_heading("6 结论", level=1)
    add_table(doc,
              ["项目", "结果", "判定"],
              [["码元速率", "4.08 MHz(FTW 定时,163 200 节拍/2M 采样精确吻合)", "通过"],
               ["占用带宽", "5.508 MHz = (1+α)Rs", "通过"],
               ["零削顶", "峰值 4094 < 4095,全程无削顶", "通过"],
               ["EVM_RMS", "4.92%(−26.2 dB),瓶颈为 ±4 码元截断,定点实现零退化", "通过(可扩展)"],
               ["频谱纯度", "噪声底 −148 dBFS/Hz,无杂散;带外 −23.6 dB 由截断主导", "通过"],
               ["可解调性", "匹配滤波 + 边界采样即恢复 PN23,眼图全开", "通过"]],
              widths=[3.4, 9.6, 3.0])

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("复现:make golden && iverilog … && vvp tb/sim.vvp +n=2000000 +amp=2797 "
                  "+dump=tb/rtl_samples.txt;python3 scripts/analyze_wave.py tb/rtl_samples.txt 2000000")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
